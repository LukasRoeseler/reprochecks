import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = r"C:\Users\lroesele.IVV5NET\Claude_Code\ReproAI\Nature Human Behavior"

with open(f"{BASE}\\NHB_DOWNLOAD_STATUS.csv", encoding="utf-8") as f:
    results = list(csv.DictReader(f))

# Load classification
with open(f"{BASE}\\NHB_MASTER_CLASSIFICATION.csv", encoding="utf-8") as f:
    classification = {r["num"]: r for r in csv.DictReader(f)}

wb = Workbook()
ws = wb.active
ws.title = "Download Status"

# Headers
headers = ['Num', 'DOI', 'DOI URL', 'First Author', 'Classification', 'Open Access', 'Full Text Available', 'Method', 'Download URL']
ws.append(headers)

# Style headers
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")

# Data
for r in results:
    cls = classification.get(r['num'], {}).get('classification', 'Unknown')
    ws.append([
        r['num'],
        r['doi'],
        r['doi_url'],
        r['first_author'],
        cls,
        'Yes' if r['is_open_access'] == 'True' else 'No',
        'Yes' if r['full_text_available'] == 'True' else 'No',
        r['method'],
        r['download_url']
    ])

# Auto-adjust column widths
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = min(max_length + 2, 60)

# Freeze panes
ws.freeze_panes = "A2"

# Summary sheet
ws2 = wb.create_sheet("Summary")
ws2.append(["Metric", "Count"])
for cell in ws2[1]:
    cell.font = Font(bold=True)

total = len(results)
oa = sum(1 for r in results if r['is_open_access'] == 'True')
available = sum(1 for r in results if r['full_text_available'] == 'True')
empirical = sum(1 for r in results if classification.get(r['num'], {}).get('classification') == 'EMPIRICAL')

ws2.append(["Total articles", total])
ws2.append(["Open Access", oa])
ws2.append(["Full text available", available])
ws2.append(["Not available", total - available])
ws2.append(["Empirical papers", empirical])

# List of available papers
ws2.append([])
ws2.append(["AVAILABLE PAPER", "DOI URL", "DOWNLOAD URL"])
for cell in ws2[ws2.max_row]:
    cell.font = Font(bold=True)

for r in results:
    if r['full_text_available'] == 'True':
        ws2.append([f"{r['num']} {r['first_author']}", r['doi_url'], r['download_url']])

# Save
output_path = f"{BASE}\\NHB_DOWNLOAD_STATUS.xlsx"
wb.save(output_path)
print(f"Created: {output_path}")
print(f"Total: {total}, Available: {available}, Not available: {total - available}")
