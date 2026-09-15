import openpyxl, sys
p = sys.argv[1]
wb = openpyxl.load_workbook(p, data_only=False)
for ws in wb.worksheets:
    print("===== SHEET:", ws.title, "dims", ws.dimensions, "maxr", ws.max_row, "maxc", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = [str(c) for c in row if c is not None and str(c).strip() != ""]
        if cells:
            print(i + 1, ":", " | ".join(cells))
