import openpyxl, io
d = r"C:\Users\lroesele.IVV5NET\AppData\Local\Temp\opencode\reprochecks\reproai_run2\09_Ohtsubo"
wb = openpyxl.load_workbook(d + r"\repro_workbook.xlsx", data_only=True)
out = io.open(d + r"\workbook_dump.txt", "w", encoding="utf-8")
for sh in wb.sheetnames:
    ws = wb[sh]
    out.write("===== SHEET: %s =====\n" % sh)
    for row in ws.iter_rows():
        vals = []
        for c in row:
            if c.value is not None:
                vals.append("%s=%r" % (c.coordinate, c.value))
        if vals:
            out.write(" | ".join(vals) + "\n")
out.close()
print("dumped")
