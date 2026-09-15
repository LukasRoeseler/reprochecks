import re, zipfile, io, sys
d = r"C:\Users\lroesele.IVV5NET\AppData\Local\Temp\opencode\reprochecks\reproai_run2\09_Ohtsubo"

# --- extract docx text ---
z = zipfile.ZipFile(d + r"\TransparencyTrail.docx")
xml = z.read("word/document.xml").decode("utf-8", "ignore")
xml = re.sub(r"</w:p>", "\n", xml)
xml = re.sub(r"<[^>]+>", "", xml)
txt = re.sub(r"\n{2,}", "\n", xml)
open(d + r"\transparencytrail.txt", "w", encoding="utf-8").write(txt)
print("docx text chars:", len(txt))

# --- xlsx ---
import openpyxl
wb = openpyxl.load_workbook(d + r"\repro_workbook.xlsx", data_only=True)
print("xlsx sheets:", wb.sheetnames)
out = io.open(d + r"\workbook_sheets.txt", "w", encoding="utf-8")
for sh in wb.sheetnames:
    ws = wb[sh]
    out.write("===== SHEET: %s  dims=%s  maxrow=%d maxcol=%d =====\n" % (sh, ws.dimensions, ws.max_row, ws.max_column))
wb2 = openpyxl.load_workbook(d + r"\repro_workbook.xlsx", data_only=False)
for sh in wb2.sheetnames:
    ws = wb2[sh]
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and isinstance(c.value, str) and c.value.startswith("="):
                out.write("%s!%s = %s\n" % (sh, c.coordinate, c.value))
out.close()
print("done")
