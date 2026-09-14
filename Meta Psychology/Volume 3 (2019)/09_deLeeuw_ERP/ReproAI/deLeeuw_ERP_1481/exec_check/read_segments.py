import openpyxl
import pandas as pd
f = r"C:\Users\lroesele.IVV5NET\Claude_Code\ReproAI\Meta Psychology\Volume 3 (2019)\09_deLeeuw_ERP\ReproAI\deLeeuw_ERP_1481\exec_check\data\generated\Good_Segments_Per_Category.xlsx"
wb = openpyxl.load_workbook(f)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("=== sheet", sn, "dims", ws.dimensions)
    for row in ws.iter_rows(values_only=True):
        print(row)
