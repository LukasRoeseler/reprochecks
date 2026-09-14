import openpyxl, pickle
fn = r'C:\Users\lroesele.IVV5NET\Claude_Code\ReproAI\Meta Psychology\Volume 4 (2020)\01_Niemeyer_PtsdPublicationBias\ReproAI\Niemeyer_PtsdPublicationBias_884\exec_check\data\Data_sets_PTSD_update_130419.xlsx'
wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
out = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    # find sei column (the effect-size standard error used by the sim: column 'sei')
    try:
        ic = header.index('sei')
    except ValueError:
        continue
    seis = [r[ic] for r in rows[1:] if r[ic] is not None and isinstance(r[ic], (int, float))]
    idv = int(name.split()[0])
    out[idv] = seis
with open(r'C:\Users\lroesele.IVV5NET\Claude_Code\ReproAI\Meta Psychology\Volume 4 (2020)\01_Niemeyer_PtsdPublicationBias\ReproAI\Niemeyer_PtsdPublicationBias_884\exec_check\output\sei_vectors.pkl', 'wb') as fh:
    pickle.dump(out, fh)
print('datasets with sei:', len(out))
print('sample:', {k: len(v) for k, v in list(out.items())[:5]})
